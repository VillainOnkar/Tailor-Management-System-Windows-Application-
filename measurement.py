from tkinter import *
from tkinter import ttk
from tkinter import messagebox
from PIL import Image,ImageTk
import mysql.connector
import os,sys
import json
import sys, os, shutil

"""def extract_google_files():
    extract_folder = os.path.join(os.getcwd(), "DriveAuth")

    # Folder auto create
    os.makedirs(extract_folder, exist_ok=True)

    # जब .exe चलता है → JSON फाइलें यहीं से मिलती हैं
    if hasattr(sys, "_MEIPASS"):
        base = sys._MEIPASS
    else:
        base = os.getcwd()

    # अब तीनों JSON file extract होंगे
    files = ["client_secrets.json", "my_drive_token.json", "config.json"]

    for file in files:
        src = os.path.join(base, file)
        dst = os.path.join(extract_folder, file)

        # अगर फाइल EXE में है और बाहर नहीं है → बाहर निकालो
        if os.path.exists(src) and not os.path.exists(dst):
            shutil.copy(src, dst)
            print("Extracted →", file)

    return extract_folder"""

from tkinter import filedialog, messagebox
APP_FOLDER_NAME = "SarthTailorShop"

BASE_PATH_FILE = os.path.join(
    os.getenv("LOCALAPPDATA"),
    APP_FOLDER_NAME,
    "base_path.json"
)

def get_sarth_base_folder():
    # Agar pehle se user ne path select kiya hai
    if os.path.exists(BASE_PATH_FILE):
        with open(BASE_PATH_FILE, "r") as f:
            return json.load(f)["base_path"]

    # FIRST TIME RUN → user se folder select karvao
    temp_root = Tk()
    temp_root.withdraw()

    messagebox.showinfo(
        "Select Folder",
        "Please select a folder where Sarth Tailor Shop data will be stored"
    )

    base_path = filedialog.askdirectory(title="Select Folder for SarthTailorShop")
    temp_root.destroy()

    if not base_path:
        sys.exit("Folder not selected")

    # Save selected path
    os.makedirs(os.path.dirname(BASE_PATH_FILE), exist_ok=True)
    with open(BASE_PATH_FILE, "w") as f:
        json.dump({"base_path": base_path}, f, indent=4)

    return base_path

def extract_google_files():
    import sys, shutil

    base_path = get_sarth_base_folder()
    extract_folder = os.path.join(
        base_path,
        APP_FOLDER_NAME,
        "DriveAuth"
    )

    os.makedirs(extract_folder, exist_ok=True)

    if hasattr(sys, "_MEIPASS"):
        src_base = sys._MEIPASS
    else:
        src_base = os.getcwd()

    files = ["client_secrets.json", "my_drive_token.json", "config.json"]

    for file in files:
        src = os.path.join(src_base, file)
        dst = os.path.join(extract_folder, file)

        if os.path.exists(src) and not os.path.exists(dst):
            shutil.copy(src, dst)

    return extract_folder


def get_connection():
    CONFIG_FILE = os.path.join(extract_google_files(), "config.json")
    with open(CONFIG_FILE, "r") as f:
        config = json.load(f)

    return mysql.connector.connect(
        host=config["host"],
        user=config["user"],
        password=config["password"],
        database=config["database"]   # yeh now tailor_db hoga
    )
 
def resource_path(relative_path):
    import sys, os
    if hasattr(sys, '_MEIPASS'):
        return os.path.join(sys._MEIPASS, relative_path)
    return os.path.join(os.path.abspath("."), relative_path)


class measurement_win:
    def __init__(self,root):
        self.root=root
        self.root.geometry("1525x535+0+262")
        self.root.title("Customer Measurement")
        
       
        self.Sh_Length_var=StringVar()
        self.Sh_Chest_var=StringVar()
        self.Sh_Stomach_var=StringVar()
        self.Sh_Shoulder_var=StringVar()
        self.Sh_Sleeves_var=StringVar()
        self.Sh_Neck_var=StringVar()
        self.Sh_Cup_var=StringVar()
        self.Sh_Front_var=StringVar()
        self.Length_var=StringVar()
        self.Weist_var=StringVar()
        self.Seat_var=StringVar()
        self.Thigh_var=StringVar()
        self.Knee_var=StringVar()
        self.Bottom_var=StringVar()
        self.shirtcutting_var=StringVar()
        self.pantcutting_var=StringVar()
        
#=======================================================lable frames==========================================
        lbl_main_title=Label(self.root,text="CUSTOMER MEASUREMENT",font=("times new roman",18,"bold"),bg="black",fg="gold",bd=4,relief=RIDGE)
        lbl_main_title.place(x=0,y=0,width=1524,height=50)

        mainframe=Frame(self.root,bd=10,relief=RIDGE)
        mainframe.place(x=0,y=50,width=1524,height=485)

        shirtframe=LabelFrame(mainframe,text="Shirt Measurement",font=("times new roman",12,"bold"),bd=10,relief=RIDGE)
        shirtframe.place(x=0,y=0,width=1505,height=200)

        pantframe=LabelFrame(mainframe,text="Pant Measurement",font=("times new roman",12,"bold"),bd=10,relief=RIDGE)
        pantframe.place(x=0,y=200,width=1505,height=200)

        thirdframe=Frame(mainframe,bd=10,relief=RIDGE)
        thirdframe.place(x=0,y=403,width=1505,height=60)
#===============================================images and btn =================================================
        imgshirt_path=resource_path(r"photos\dress1.png")
        imgshirt=Image.open(imgshirt_path)
        imgshirt=imgshirt.resize((170,170),Image.ANTIALIAS)
        self.imgshirt_demo=ImageTk.PhotoImage(imgshirt)
        imgshirt_lbl=Label(shirtframe,image=self.imgshirt_demo,bd=0,relief=RIDGE)
        imgshirt_lbl.place(x=0,y=0,width=170,height=170)

        imgpant_path=resource_path(r"photos\pant.png")
        imgpant=Image.open(imgpant_path)
        imgpant=imgpant.resize((125,173),Image.ANTIALIAS)
        self.imgpant_demo=ImageTk.PhotoImage(imgpant)
        imgpant_lbl=Label(pantframe,image=self.imgpant_demo,bd=0,relief=RIDGE)
        imgpant_lbl.place(x=1320,y=0,width=125,height=173)

        imgrightcolor_path=resource_path(r"photos\rightcolor.png")
        imgrightcolor=Image.open(imgrightcolor_path)
        imgrightcolor=imgrightcolor.resize((470,50),Image.ANTIALIAS)
        self.imgrightcolor_demo=ImageTk.PhotoImage(imgrightcolor)
        imgrightcolor_lbl=Label(thirdframe,image=self.imgrightcolor_demo,bd=0,relief=RIDGE)
        imgrightcolor_lbl.place(x=1015,y=0,width=470,height=40)

        imgleftcolor_path=resource_path(r"photos\leftcolor.png")
        imgleftcolor=Image.open(imgleftcolor_path)
        imgleftcolor=imgleftcolor.resize((400,50),Image.ANTIALIAS)
        self.imgleftcolor_demo=ImageTk.PhotoImage(imgleftcolor)
        imgleftcolor_lbl=Label(thirdframe,image=self.imgleftcolor_demo,bd=0,relief=RIDGE)
        imgleftcolor_lbl.place(x=0,y=0,width=400,height=40)

        imgcentercolor_path=resource_path(r"photos\leftcolor.png")
        imgcentercolor=Image.open(imgcentercolor_path)
        imgcentercolor=imgleftcolor.resize((400,50),Image.ANTIALIAS)
        self.imgcentercolor_demo=ImageTk.PhotoImage(imgcentercolor)
        imgcentercolor_lbl=Label(thirdframe,image=self.imgcentercolor_demo,bd=0,relief=RIDGE)
        imgcentercolor_lbl.place(x=550,y=0,width=400,height=40)

        imgcutpiecewhite_path=resource_path(r"photos\white.png")
        imgcutpiecewhite=Image.open(imgcutpiecewhite_path)
        imgcutpiecewhite= imgcutpiecewhite.resize((200,150),Image.ANTIALIAS)
        self.imgcutpiecewhite_demo=ImageTk.PhotoImage(imgcutpiecewhite)
        imgcutpiecewhite_lbl=Label(pantframe,image=self.imgcutpiecewhite_demo,bd=0,relief=RIDGE)
        imgcutpiecewhite_lbl.place(x=0,y=10,width=200,height=150)

        imgcutpieceblue_path=resource_path(r"photos\blue.png")
        imgcutpieceblue=Image.open(imgcutpieceblue_path)
        imgcutpieceblue=imgcutpieceblue.resize((200,150),Image.ANTIALIAS)
        self.imgcutpieceblue_demo=ImageTk.PhotoImage(imgcutpieceblue)
        imgcutpieceblue_lbl=Label(pantframe,image=self.imgcutpieceblue_demo,bd=0,relief=RIDGE)
        imgcutpieceblue_lbl.place(x=200,y=10,width=200,height=150)

        imgmixcolor_path=resource_path(r"photos\color.png")
        imgmixcolor=Image.open(imgmixcolor_path)
        imgmixcolor=imgmixcolor.resize((280,150),Image.ANTIALIAS)
        self.imgmixcolor_demo=ImageTk.PhotoImage(imgmixcolor)
        imgmixcolor_lbl=Label(shirtframe,image=self.imgmixcolor_demo,bd=0,relief=RIDGE)
        imgmixcolor_lbl.place(x=1200,y=10,width=280,height=150)
#============================================================= search field and button======================================
#=============================================================variables===========================================
        self.name1_var=StringVar()
#=================================================================================================================
        name_lbl_=Label(thirdframe,text="NAME",font=("times new roman",15,"bold"),bg="red",width=11,bd=4)
        name_lbl_.place(x=466,y=5)
        """conn=mysql.connector.connect(
                                               host="localhost",
                                               user="root",
                                               password="OnkarA@060105",
                                               database="sarthtailor"
                                               )
        mycursor=conn.cursor()
        mycursor.execute("select Name from addcustomer")
        rows=mycursor.fetchall()

        name_combobox=ttk.Combobox(thirdframe,textvariable=self.name1_var,font=("times new roman",12,"bold"),width=20,state="readonly")
        name_combobox["values"]=rows
        name_combobox.current(0)
        name_combobox.place(x=620,y=8)"""
        namerntry=ttk.Entry(thirdframe,textvariable=self.name1_var,font=("times new roman",12,"bold"),width=20)
        namerntry.place(x=620,y=8)

        #addbtn
        addbtn=Button(thirdframe,text="ADD",command=self.addmeasurement_btn,font=("times new roman",15,"bold"),bd=4,relief=RIDGE,bg="green",fg="black",activebackground="green",activeforeground="black",cursor="hand2")
        addbtn.place(x=800,y=5,width=150,height=30)

        #updatebtn
        updatebtn=Button(thirdframe,text="UPDATE",command=self.updatedbmsbtn,font=("times new roman",15,"bold"),bd=4,relief=RIDGE,bg="orange",fg="black",activebackground="orange",activeforeground="black",cursor="hand2")
        updatebtn.place(x=960,y=5,width=150,height=30)

        #resetbtn
        resetbtn=Button(thirdframe,text="RESET",command=self.reset_btn,font=("times new roman",15,"bold"),bd=4,relief=RIDGE,bg="cyan",fg="black",activebackground="cyan",activeforeground="black",cursor="hand2")
        resetbtn.place(x=305,y=5,width=150,height=30)
#=================================================shirt and pant labals and entry fields============================================
        #shirt
        sh_length_lbl=Label(shirtframe,text="Sh_Length",font=("times new roman",12,"bold"))
        sh_length_lbl.place(x=180,y=20)
        sh_lengthentryentry=ttk.Entry(shirtframe,textvariable=self.Sh_Length_var,font=("times new roman",12,"bold"))
        sh_lengthentryentry.place(x=265,y=20)

        shchest=Label(shirtframe,text="Sh_Chest",font=("times new roman",12,"bold"))
        shchest.place(x=480,y=20)
        shchestentry=ttk.Entry(shirtframe,textvariable=self.Sh_Chest_var,font=("times new roman",12,"bold"))
        shchestentry.place(x=550,y=20)

        shstomach=Label(shirtframe,text="Sh_Stomach",font=("times new roman",12,"bold"))
        shstomach.place(x=780,y=20)
        shstomachentry=ttk.Entry(shirtframe,textvariable=self.Sh_Stomach_var,font=("times new roman",12,"bold"))
        shstomachentry.place(x=865,y=20)

        shshoulder=Label(shirtframe,text="Sh_Shoulder",font=("times new roman",12,"bold"))
        shshoulder.place(x=1080,y=20)
        shshoulderentry=ttk.Entry(shirtframe,textvariable=self.Sh_Shoulder_var,font=("times new roman",12,"bold"))
        shshoulderentry.place(x=1170,y=20)

        shsleeves=Label(shirtframe,text="Sh_Sleeves",font=("times new roman",12,"bold"))
        shsleeves.place(x=180,y=100)
        shsleevesentry=ttk.Entry(shirtframe,textvariable=self.Sh_Sleeves_var,font=("times new roman",12,"bold"))
        shsleevesentry.place(x=265,y=100)

        shneck=Label(shirtframe,text="Sh_Neck",font=("times new roman",12,"bold"))
        shneck.place(x=480,y=100)
        shneckentry=ttk.Entry(shirtframe,textvariable=self.Sh_Neck_var,font=("times new roman",12,"bold"))
        shneckentry.place(x=550,y=100)

        shcup=Label(shirtframe,text="Sh_Cup",font=("times new roman",12,"bold"))
        shcup.place(x=780,y=100)
        shcupentry=ttk.Entry(shirtframe,textvariable=self.Sh_Cup_var,font=("times new roman",12,"bold"))
        shcupentry.place(x=850,y=100)

        shfront=Label(shirtframe,text="Sh_Front",font=("times new roman",12,"bold"))
        shfront.place(x=1080,y=100)
        shfrontentry=ttk.Entry(shirtframe,textvariable=self.Sh_Front_var,font=("times new roman",12,"bold"))
        shfrontentry.place(x=1150,y=100)

        lbl_shirtcutting=Label(shirtframe,text="Shirt Cutting:",font=("times new roman",12,"bold"))
        lbl_shirtcutting.place(x=625,y=135)
        
        yes_btn1=Radiobutton(shirtframe,variable=self.shirtcutting_var,value="Yes",text="Yes",font=("times new roman",13,"bold"))
        yes_btn1.place(x=720,y=135)
        self.shirtcutting_var.set(value="No")
        
        no_btn1=Radiobutton(shirtframe,variable=self.shirtcutting_var,value="No",text="No",font=("times new roman",13,"bold"))
        no_btn1.place(x=780,y=135)

        #pant 
        length=Label(pantframe,text="Length",font=("times new roman",12,"bold"))
        length.place(x=480,y=20)
        lengthentry=ttk.Entry(pantframe,textvariable=self.Length_var,font=("times new roman",12,"bold"))
        lengthentry.place(x=550,y=20)

        weist=Label(pantframe,text="Weist",font=("times new roman",12,"bold"))
        weist.place(x=780,y=20)
        weistentry=ttk.Entry(pantframe,textvariable=self.Weist_var,font=("times new roman",12,"bold"))
        weistentry.place(x=850,y=20)

        seat=Label(pantframe,text="Seat",font=("times new roman",12,"bold"))
        seat.place(x=1080,y=20)
        seatentry=ttk.Entry(pantframe,textvariable=self.Seat_var,font=("times new roman",12,"bold"))
        seatentry.place(x=1150,y=20)

        thigh=Label(pantframe,text="Thigh",font=("times new roman",12,"bold"))
        thigh.place(x=480,y=100)
        thighentry=ttk.Entry(pantframe,textvariable=self.Thigh_var,font=("times new roman",12,"bold"))
        thighentry.place(x=550,y=100)

        knee=Label(pantframe,text="Knee",font=("times new roman",12,"bold"))
        knee.place(x=780,y=100)
        kneeentry=ttk.Entry(pantframe,textvariable=self.Knee_var,font=("times new roman",12,"bold"))
        kneeentry.place(x=850,y=100)

        bottom=Label(pantframe,text="Bottom",font=("times new roman",12,"bold"))
        bottom.place(x=1080,y=100)
        bottomentry=ttk.Entry(pantframe,textvariable=self.Bottom_var,font=("times new roman",12,"bold"))
        bottomentry.place(x=1150,y=100)

        lbl_shirtcutting=Label(pantframe,text="Pant Cutting:",font=("times new roman",12,"bold"))
        lbl_shirtcutting.place(x=635,y=135)
        
        yes_btn=Radiobutton(pantframe,variable=self.pantcutting_var,value="Yes",text="Yes",font=("times new roman",13,"bold"))
        yes_btn.place(x=730,y=135)
        self.pantcutting_var.set(value="No")
        
        no_btn=Radiobutton(pantframe,variable=self.pantcutting_var,value="No",text="No",font=("times new roman",13,"bold"))
        no_btn.place(x=790,y=135)


######################################################database gui###################################################################
        lbl_main_title=Label(self.root,text="DETAILS",font=("times new roman",18,"bold"),bg="black",fg="gold",bd=4,relief=RIDGE)
        lbl_main_title.place(x=0,y=530,width=1524,height=30)

        mainframe=Frame(self.root,bd=10,relief=RIDGE)
        mainframe.place(x=0,y=560,width=1524,height=280)
#==================================================photos======================================================
        
        imgcolor2_path=resource_path("photos\color2.png")
        imgcolor2=Image.open(imgcolor2_path)
        imgcolor2=imgcolor2.resize((470,472),Image.ANTIALIAS)
        self.imgcolor2_demo=ImageTk.PhotoImage(imgcolor2)
        imgcolor2_lbl=Label(mainframe,image=self.imgcolor2_demo,bd=0,relief=RIDGE)
        imgcolor2_lbl.place(x=1039,y=0,width=470,height=265)

#==========================================================serching system=========================================
#======================================================variables==========================
        self.search_combobox_var=StringVar()
        self. search_entry_var=StringVar()

        lbl_search=Label(mainframe,text="SEARCH BY",font=("times new roman",12,"bold"),bg="red",fg="white")
        lbl_search.place(x=1110,y=20)

        search_combobox=ttk.Combobox(mainframe,textvariable=self.search_combobox_var,font=("times new roman",12,"bold"),width=15,state="readonly")
        search_combobox["values"]=("Select Option","Name","Shirt_Cutting","Pant_Cutting")
        search_combobox.current(0)
        search_combobox.place(x=1080,y=50)

        search_entry=Entry(mainframe,textvariable=self.search_entry_var,font=("times new roman",13,"bold"),width=15)
        search_entry.place(x=1080,y=80)

        search_btn=Button(mainframe,text="SEARCH",command=self.search_btn,font=("times new roman",12,"bold"),bd=4,bg="cyan",fg="black",width=15,activebackground="cyan",activeforeground="black",cursor="hand2")
        search_btn.place(x=1080,y=110)

        deletebtn=Button(mainframe,text="DELETE",command=self.deletedbms_btn,font=("times new roman",12,"bold"),bd=4,bg="red",fg="white",width=15,activebackground="red",activeforeground="white",cursor="hand2")
        deletebtn.place(x=1080,y=150)

        showallbtn=Button(mainframe,text="SHOW All",command=self.fecth_data,font=("times new roman",12,"bold"),bd=4,bg="black",fg="gold",width=15,activebackground="black",activeforeground="gold",cursor="hand2")
        showallbtn.place(x=1080,y=190)

        resetbtn2=Button(mainframe,text="RESET",command=self.reset_btn2,font=("times new roman",15,"bold"),bd=4,relief=RIDGE,bg="cyan",fg="black",activebackground="cyan",activeforeground="black",cursor="hand2")
        resetbtn2.place(x=1080,y=230,width=150,height=30)

        downloaddatabtn=Button(mainframe,text="DOWNLOAD RECORD",command=self.export_to_word_measurement,font=("times new roman",10,"bold"),bd=4,relief=RIDGE,bg="black",fg="gold",activebackground="black",activeforeground="gold",cursor="hand2")
        downloaddatabtn.place(x=1355,y=0,width=150,height=30)

#=============================================================tables================================================
        databaseframe=Frame(mainframe,bd=10,relief=RIDGE)
        databaseframe.place(x=0,y=0,width=1042,height=262)

        scroll_x2=Scrollbar(databaseframe,orient=HORIZONTAL)
        scroll_x2.pack(side=BOTTOM,fill=X)

        scroll_y2=Scrollbar(databaseframe,orient=VERTICAL)
        scroll_y2.pack(side=RIGHT,fill=Y)

        self.databaseframe=ttk.Treeview(databaseframe,
                                    columns=( 
                                              "Name",
                                              "Shirt_Cutting",
                                              "Pant_Cutting",
                                              "Sh_Length",
                                              "Sh_Chest",
                                              "Sh_Stomach",
                                              "Sh_Shoulder",
                                              "Sh_Sleeves",
                                              "Sh_Neck",
                                              "Sh_Cup",
                                              "Sh_Front",
                                              "Length",
                                              "Waist",
                                              "Seat",
                                              "Thigh",
                                              "Knee",
                                              "Bottom",
                                               ),xscrollcommand=scroll_x2.set,yscrollcommand=scroll_y2.set)
        
        scroll_x2.pack(side=BOTTOM,fill=X)
        scroll_y2.pack(side=RIGHT,fill=Y)

        scroll_x2.config(command=self.databaseframe.xview)
        scroll_y2.config(command=self.databaseframe.yview)

        self.databaseframe["show"]="headings"

        self.databaseframe.heading("Name",text="Name")
        self.databaseframe.heading("Shirt_Cutting",text="Shirt_Cutting")
        self.databaseframe.heading("Pant_Cutting",text="Pant_Cutting")
        self.databaseframe.heading("Sh_Length",text="Sh_Length")
        self.databaseframe.heading("Sh_Chest",text="Sh_Chest")
        self.databaseframe.heading("Sh_Stomach",text="Sh_Stomach")
        self.databaseframe.heading("Sh_Shoulder",text="Sh_Shoulder")
        self.databaseframe.heading("Sh_Sleeves",text="Sh_Sleeves")
        self.databaseframe.heading("Sh_Neck",text="Sh_Neck")
        self.databaseframe.heading("Sh_Cup",text="Sh_Cup")
        self.databaseframe.heading("Sh_Front",text="Sh_Front")
        self.databaseframe.heading("Length",text="Length")
        self.databaseframe.heading("Waist",text="Waist")
        self.databaseframe.heading("Seat",text="Seat")
        self.databaseframe.heading("Thigh",text="Thigh")
        self.databaseframe.heading("Knee",text="Knee")
        self.databaseframe.heading("Bottom",text="Bottom")
        self.databaseframe.pack(fill=BOTH,expand=1)

        self.databaseframe.column("Name",width=100)
        self.databaseframe.column("Shirt_Cutting",width=100)
        self.databaseframe.column("Pant_Cutting",width=100)
        self.databaseframe.column("Sh_Length",width=100)
        self.databaseframe.column("Sh_Chest",width=100)
        self.databaseframe.column("Sh_Stomach",width=100)
        self.databaseframe.column("Sh_Shoulder",width=100)
        self.databaseframe.column("Sh_Sleeves",width=100)
        self.databaseframe.column("Sh_Neck",width=100)
        self.databaseframe.column("Sh_Cup",width=100)
        self.databaseframe.column("Sh_Front",width=100)
        self.databaseframe.column("Length",width=100)
        self.databaseframe.column("Waist",width=100)
        self.databaseframe.column("Seat",width=100)
        self.databaseframe.column("Thigh",width=100)
        self.databaseframe.column("Knee",width=100)
        self.databaseframe.column("Bottom",width=100)
        self.fecth_data()
        self.databaseframe.bind("<ButtonRelease-1>",self.getcursor)
                
#=========================================================database connectivity==================================================
        self.auto_export_to_csv_measurement()
    def addmeasurement_btn(self):
        if self.name1_var.get()=="":
             messagebox.showerror("Error","Fill The Name Box Its The Compulsary Box...!",parent=self.root)
        else:
             conn=get_connection()
             mycursor=conn.cursor()
             mycursor.execute("insert into measurement values(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)",(
                                                                                                         self.name1_var.get(),
                                                                                                         self.shirtcutting_var.get(),
                                                                                                         self.pantcutting_var.get(),
                                                                                                         self.Sh_Length_var.get(),
                                                                                                         self.Sh_Chest_var.get(),
                                                                                                         self.Sh_Stomach_var.get(),
                                                                                                         self.Sh_Shoulder_var.get(),
                                                                                                         self.Sh_Sleeves_var.get(),
                                                                                                         self.Sh_Neck_var.get(),
                                                                                                         self.Sh_Cup_var.get(),
                                                                                                         self.Sh_Front_var.get(),
                                                                                                         self.Length_var.get(),
                                                                                                         self.Weist_var.get(),
                                                                                                         self.Seat_var.get(),
                                                                                                         self.Thigh_var.get(),
                                                                                                         self.Knee_var.get(),
                                                                                                         self.Bottom_var.get(),
                                                                                                         

                                                                                                         ))
        conn.commit()
        self.fecth_data()
        messagebox.showinfo("Success",f"Customer {self.name1_var.get()} Measurement Has Been Added Successfully...!",parent=self.root)
        conn.close()



    def fecth_data(self):
        conn=get_connection()
        mycursor=conn.cursor()
        mycursor.execute("select * from measurement")
        rows=mycursor.fetchall()
        if len(rows)!=0:
                self.databaseframe.delete(*self.databaseframe.get_children())
                for i in rows:
                     self.databaseframe.insert("",END,values=i)
                conn.commit()
        conn.close()


    def getcursor(self,event=""):
        cursor=self.databaseframe.focus()
        content=self.databaseframe.item(cursor)
        rows=content["values"]

        self.name1_var.set(rows[0])
        self.shirtcutting_var.set(rows[1])
        self.pantcutting_var.set(rows[2])
        self.Sh_Length_var.set(rows[3]),
        self.Sh_Chest_var.set(rows[4]),
        self.Sh_Stomach_var.set(rows[5]),
        self.Sh_Shoulder_var.set(rows[6]),
        self.Sh_Sleeves_var.set(rows[7]),
        self.Sh_Neck_var.set(rows[8]),
        self.Sh_Cup_var.set(rows[9]),
        self.Sh_Front_var.set(rows[10]),
        self.Length_var.set(rows[11]),
        self.Weist_var.set(rows[12]),
        self.Seat_var.set(rows[13]),
        self.Thigh_var.set(rows[14]),
        self.Knee_var.set(rows[15]),
        self.Bottom_var.set(rows[16]),

    def updatedbmsbtn(self):
        conn=get_connection()
        mycursor=conn.cursor()
        mycursor.execute("update measurement set Shirt_Cutting=%s,Pant_Cutting=%s,Sh_Length=%s,Sh_Chest=%s,Sh_Stomach=%s,Sh_Shoulder=%s,Sh_Sleeves=%s,Sh_Neck=%s,Sh_Cup=%s,Sh_Front=%s,Length=%s,Waist=%s,Seat=%s,Thigh=%s,Knee=%s,Bottom=%s where Name=%s",
                                                                                             (
                                                                                                self.shirtcutting_var.get(),
                                                                                                self.pantcutting_var.get(),
                                                                                                self.Sh_Length_var.get(),
                                                                                                self.Sh_Chest_var.get(),
                                                                                                self.Sh_Stomach_var.get(),
                                                                                                self.Sh_Shoulder_var.get(),
                                                                                                self.Sh_Sleeves_var.get(),
                                                                                                self.Sh_Neck_var.get(),
                                                                                                self.Sh_Cup_var.get(),
                                                                                                self.Sh_Front_var.get(),
                                                                                                self.Length_var.get(),
                                                                                                self.Weist_var.get(),
                                                                                                self.Seat_var.get(),
                                                                                                self.Thigh_var.get(),
                                                                                                self.Knee_var.get(),
                                                                                                self.Bottom_var.get(),
                                                                                                self.name1_var.get()

                                                                                        
                                                                                                
                                                                                                ))
        conn.commit()
        self.fecth_data()
        conn.close()
        messagebox.showinfo("Success",f"Customer {self.name1_var.get()} Measurement Has Been Update Successfully...! ",parent=self.root)



    def deletedbms_btn(self):
        mdelete=messagebox.askyesno("Sarth Tailor","Do you want delete this customer",parent=self.root)
        if mdelete>0:
             conn=get_connection()
             mycursor=conn.cursor()
             query="delete from measurement where Name=%s"
             val=(self.name1_var.get(),)
             mycursor.execute(query,val)
        else:
                if not mdelete:
                        return

        conn.commit()
        self.fecth_data()
        conn.close()
        messagebox.showinfo("Success",f"Customer {self.name1_var.get()} Record Has Been Deleted...!",parent=self.root)



    def reset_btn(self):
        self.name1_var.set("")
        self.shirtcutting_var.set("No")
        self.pantcutting_var.set("No")
        self.Sh_Length_var.set(""),
        self.Sh_Chest_var.set(""),
        self.Sh_Stomach_var.set(""),
        self.Sh_Shoulder_var.set(""),
        self.Sh_Sleeves_var.set(""),
        self.Sh_Neck_var.set(""),
        self.Sh_Cup_var.set(""),
        self.Sh_Front_var.set(""),
        self.Length_var.set(""),
        self.Weist_var.set(""),
        self.Seat_var.set(""),
        self.Thigh_var.set(""),
        self.Knee_var.set(""),
        self.Bottom_var.set(""),

    def search_btn(self):
        if self.search_combobox_var.get()=="Select Option" or self.search_entry_var.get()=="":
            messagebox.showerror("Error","All Fields Are Required",parent=self.root)
        else:
            conn=get_connection()
            mycursor=conn.cursor()
            mycursor.execute("SELECT * FROM measurement WHERE " + str(self.search_combobox_var.get()) + " LIKE '%" + str(self.search_entry_var.get()) + "%'")
            rows=mycursor.fetchall()
            if len(rows)!=0:
                self.databaseframe.delete(*self.databaseframe.get_children())
                for i in rows:
                    self.databaseframe.insert("",END,values=i)
                conn.commit()
            conn.close()


    def reset_btn2(self):
         self.search_combobox_var.set("Select Option")
         self.search_entry_var.set("")


    def auto_export_to_csv_measurement(self):
        import csv
        import mysql.connector
        import hashlib

        try:
            conn = get_connection()
            cursor = conn.cursor()
            cursor.execute("SELECT * FROM measurement ORDER BY Name ASC")
            all_data = cursor.fetchall()
            conn.close()

            # ---- Make a hash of current DB state ----
            db_hash = hashlib.md5(str(all_data).encode()).hexdigest()

            # ---- If previous hash exists, compare ----
            if hasattr(self, "last_db_hash") and self.last_db_hash == db_hash:
                print("Success","⏸ No change detected — CSV not updated.")
            else:
                self.last_db_hash = db_hash  # Save new hash

                headers = [
                    "Name","Shirt_Cutting","Pant_Cutting","Sh_Length","Sh_Chest","Sh_Stomach","Sh_Shoulder","Sh_Sleeves","Sh_Neck","Sh_Cup","Sh_Front",
                    "Length","Weist","Seat","Thigh","Knee","Bottom"
                ]

                #file_path = r"E:\Tailor management system\measurement_data.csv"
                file_path = resource_path("measurement_data.csv")

                with open(file_path, mode="w", newline="", encoding="utf-8") as file:
                    writer = csv.writer(file)
                    writer.writerow(headers)
                    writer.writerows(all_data)

                print("Success","CSV updated (change detected).")

        except Exception as e:
            print(f"Error updating CSV: {e}")

        # 🔁 Repeat every 5 seconds
        self.root.after(5000, self.auto_export_to_csv_measurement)

    def export_to_word_measurement(self):
        import pandas as pd
        from docx import Document
        from tkinter import filedialog, messagebox

        # Step 1: Confirmation popup
        confirm = messagebox.askyesno("Confirmation", "Do you really want to export customer measurement data to a Word file?",parent=self.root)
        if not confirm:
            return  # Agar user "No" dabaye to function yahi ruk jaye

        try:
            #csv_path = "E:\\Tailor management system\\measurement_data.csv"
            csv_path = resource_path("measurement_data.csv")
            df = pd.read_csv(csv_path)

            #  Step 2: Save file dialog
            save_path = filedialog.asksaveasfilename(
                defaultextension=".docx",
                filetypes=[("Word Documents", "*.docx")],
                title="Save Word File As"
            )

            if not save_path:
                return  # User ne cancel kiya

            #  Step 3: Word document create karo
            doc = Document()
            doc.add_heading("Customer Measurement Details", level=1)

            #  Step 4: Table banao aur header add karo
            table = doc.add_table(rows=1, cols=len(df.columns))
            hdr_cells = table.rows[0].cells
            for i, col_name in enumerate(df.columns):
                hdr_cells[i].text = col_name

            #  Step 5: Data rows add karo
            for _, row in df.iterrows():
                row_cells = table.add_row().cells
                for i, value in enumerate(row):
                    row_cells[i].text = str(value)

            # Step 6: Save Word file
            doc.save(save_path)
            messagebox.showinfo("Success", f"Word file successfully saved at:\n{save_path}", parent=self.root)

        except Exception as e:
            messagebox.showerror("Error", f"Failed to export: {str(e)}", parent=self.root)

    



if __name__=="__main__":
    root=Tk()
    app=measurement_win(root)
    root.mainloop()
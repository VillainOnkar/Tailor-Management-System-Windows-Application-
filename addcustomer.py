from tkinter import *
from tkinter import ttk
from tkinter import messagebox
from PIL import Image,ImageTk
import mysql.connector
import random 
from tkcalendar import DateEntry
import requests
import json
import os

import sys, os, shutil

"""def extract_google_files():
    extract_folder = os.path.join(os.getcwd(), "DriveAuth")

    # Folder auto create
    os.makedirs(extract_folder, exist_ok=True)

    # जब .exe चलता है → JSON फाइलें यहीं से मिलती हैं
    if hasattr(sys, "_MEIPASS"):
        base = sys._MEIPASS
    else:
        base = os.getcwd()

    # अब तीनों JSON file extract होंगे
    files = ["client_secrets.json", "my_drive_token.json", "config.json"]

    for file in files:
        src = os.path.join(base, file)
        dst = os.path.join(extract_folder, file)

        # अगर फाइल EXE में है और बाहर नहीं है → बाहर निकालो
        if os.path.exists(src) and not os.path.exists(dst):
            shutil.copy(src, dst)
            print("Extracted →", file)

    return extract_folder"""
from tkinter import filedialog, messagebox
APP_FOLDER_NAME = "SarthTailorShop"

BASE_PATH_FILE = os.path.join(
    os.getenv("LOCALAPPDATA"),
    APP_FOLDER_NAME,
    "base_path.json"
)

def get_sarth_base_folder():
    # Agar pehle se user ne path select kiya hai
    if os.path.exists(BASE_PATH_FILE):
        with open(BASE_PATH_FILE, "r") as f:
            return json.load(f)["base_path"]

    # FIRST TIME RUN → user se folder select karvao
    temp_root = Tk()
    temp_root.withdraw()

    messagebox.showinfo(
        "Select Folder",
        "Please select a folder where Sarth Tailor Shop data will be stored"
    )

    base_path = filedialog.askdirectory(title="Select Folder for SarthTailorShop")
    temp_root.destroy()

    if not base_path:
        sys.exit("Folder not selected")

    # Save selected path
    os.makedirs(os.path.dirname(BASE_PATH_FILE), exist_ok=True)
    with open(BASE_PATH_FILE, "w") as f:
        json.dump({"base_path": base_path}, f, indent=4)

    return base_path

def extract_google_files():
    import sys, shutil

    base_path = get_sarth_base_folder()
    extract_folder = os.path.join(
        base_path,
        APP_FOLDER_NAME,
        "DriveAuth"
    )

    os.makedirs(extract_folder, exist_ok=True)

    if hasattr(sys, "_MEIPASS"):
        src_base = sys._MEIPASS
    else:
        src_base = os.getcwd()

    files = ["client_secrets.json", "my_drive_token.json", "config.json"]

    for file in files:
        src = os.path.join(src_base, file)
        dst = os.path.join(extract_folder, file)

        if os.path.exists(src) and not os.path.exists(dst):
            shutil.copy(src, dst)

    return extract_folder


def get_connection():
        CONFIG_FILE = os.path.join(extract_google_files(), "config.json")
        with open(CONFIG_FILE, "r") as f:
            config = json.load(f)

        return mysql.connector.connect(
            host=config["host"],
            user=config["user"],
            password=config["password"],
            database=config["database"]   # yeh now tailor_db hoga
        )


def resource_path(relative_path):
    import sys, os
    if hasattr(sys, '_MEIPASS'):
        return os.path.join(sys._MEIPASS, relative_path)
    return os.path.join(os.path.abspath("."), relative_path)


class customer:
    def __init__(self,root):
        self.root=root
        self.root.geometry("1298x538+227+263")
        self.root.title("Adding New Customer")
        self.root.resizable(False,False)

#================================================variables===========================================
        self.reference_var=StringVar()
        x=random.randint(1000,9999)
        self.reference_var.set(str(x))

        self.name=StringVar()
        self.address=StringVar()
        self.phone=StringVar()
        self.Shirtquantity=StringVar()
        self.Pantquantity=StringVar()
        self.type_of_shirt=StringVar()
        self.design_of_shirt=StringVar()
        self.receiving_date=StringVar()
        self.submitted_date=StringVar()
        self.payment=StringVar()
        self.delivery_var=StringVar()
        
#========================================labals enrty =======================================
        lbl_main_title=Label(self.root,text="ADD CUSTOMER INFO",font=("times new roman",18,"bold"),bg="black",fg="gold",bd=4,relief=RIDGE)
        lbl_main_title.place(x=0,y=0,width=1297,height=50)

        """img_logo_path=resource_path(r"E:\Tailor management system\photos\logo.jpeg")
        img_logo=Image.open(img_logo_path)
        img_logo=img_logo.resize((100,50),Image.ANTIALIAS)
        self.logodemo=ImageTk.PhotoImage(img_logo)
        lbl_img_logo=Label(self.root,image=self.logodemo,bd=0,relief=RIDGE)
        lbl_img_logo.place(x=5,y=5,width=100,height=40)"""

        img_logo_path=resource_path("photos\logo.jpeg")
        img_logo=Image.open(img_logo_path)
        img_logo=img_logo.resize((100,50),Image.ANTIALIAS)
        self.logodemo=ImageTk.PhotoImage(img_logo)
        lbl_img_logo=Label(self.root,image=self.logodemo,bd=0,relief=RIDGE)
        lbl_img_logo.place(x=5,y=5,width=100,height=40)
       
        left_frame=LabelFrame(self.root,bd=10,relief=RIDGE,text="Customer Details",font=("times new roman",12,"bold"),padx=2)
        left_frame.place(x=5,y=50,width=425,height=480)

        #referenceno
        referenceno_lbl=Label(left_frame,text="Receipt No",font=("times new roman",12,"bold"))
        referenceno_lbl.grid(row=1,column=0,sticky=W,pady=5)
        referenceno_entry= ttk.Entry(left_frame,textvariable=self.reference_var,font=("times new roman",12,"bold"),width=29,state="readonly")
        referenceno_entry.grid(row=1,column=1,pady=5)
       
        #Name 
        name_lbl=Label(left_frame,text="Name:",font=("times new roman",12,"bold"))
        name_lbl.grid(row=2,column=0,sticky=W,pady=5)
        name_entry= ttk.Entry(left_frame,textvariable=self.name,font=("times new roman",12,"bold"),width=29)
        name_entry.grid(row=2,column=1,pady=5)

        #Address
        name_lbl=Label(left_frame,text="Address:",font=("times new roman",12,"bold"))
        name_lbl.grid(row=3,column=0,sticky=W,pady=5)
        name_entry= ttk.Entry(left_frame,textvariable=self.address,font=("times new roman",12,"bold"),width=29)
        name_entry.grid(row=3,column=1,pady=5)

        #phone
        name_lbl=Label(left_frame,text="Phone:",font=("times new roman",12,"bold"))
        name_lbl.grid(row=4,column=0,sticky=W,pady=5)
        name_entry= ttk.Entry(left_frame,textvariable=self.phone,font=("times new roman",12,"bold"),width=29)
        name_entry.grid(row=4,column=1,pady=5)

        #Shirt Quantity
        name_lbl=Label(left_frame,text="Shirt Quantity:",font=("times new roman",12,"bold"))
        name_lbl.grid(row=5,column=0,sticky=W,pady=5)
        name_entry= ttk.Entry(left_frame,textvariable=self.Shirtquantity,font=("times new roman",12,"bold"),width=29)
        name_entry.grid(row=5,column=1,pady=5)

        #Pant Quantity
        name_lbl=Label(left_frame,text="Pant Quantity:",font=("times new roman",12,"bold"))
        name_lbl.grid(row=6,column=0,sticky=W,pady=5)
        name_entry= ttk.Entry(left_frame,textvariable=self.Pantquantity,font=("times new roman",12,"bold"),width=29)
        name_entry.grid(row=6,column=1,pady=5)

        # Type of shirt
        name_lbl=Label(left_frame,text="Type Of Shirt:",font=("times new roman",12,"bold"))
        name_lbl.grid(row=7,column=0,sticky=W,pady=5)
        typeofshirt_combobox=ttk.Combobox(left_frame,textvariable=self.type_of_shirt,font=("times new roman",12,"bold"),state="readonly",width=27)
        typeofshirt_combobox['values']=("Select Option","Full Shirt","Half Shirt","Kurta")
        typeofshirt_combobox.current(0)
        typeofshirt_combobox.grid(row=7,column=1)

        #design of shirt
        designofshirt_lbl=Label(left_frame,text="Design Of Shirt",font=("times new roman",12,"bold"))
        designofshirt_lbl.grid(row=8,column=0,sticky=W,pady=5)
        designofshirt_combobox=ttk.Combobox(left_frame,textvariable=self.design_of_shirt,font=("times new roman",12,"bold"),state="readonly",width=27)
        designofshirt_combobox["values"]=("Select Option","Normal Cut","Apple Cut")
        designofshirt_combobox.current(0)
        designofshirt_combobox.grid(row=8,column=1)

        #receiving date
        name_lbl=Label(left_frame,text="Receiving Date:",font=("times new roman",12,"bold"))
        name_lbl.grid(row=9,column=0,sticky=W,pady=5)
        name_entry= DateEntry(left_frame,textvariable=self.receiving_date,font=("times new roman",12,"bold"),width=26,date_pattern="dd-mm-yyyy")
        name_entry.grid(row=9,column=1,pady=5)

        #submitted date
        name_lbl=Label(left_frame,text="Submitted Date:",font=("times new roman",12,"bold"))
        name_lbl.grid(row=10,column=0,sticky=W,pady=5)
        name_entry= DateEntry(left_frame,textvariable=self.submitted_date,font=("times new roman",12,"bold"),width=26,date_pattern="dd-mm-yyyy")
        name_entry.grid(row=10,column=1,pady=5)

        #payment
        name_lbl=Label(left_frame,text="Payment:",font=("times new roman",12,"bold"))
        name_lbl.grid(row=11,column=0,sticky=W,pady=5)
        name_entry= ttk.Entry(left_frame,textvariable=self.payment,font=("times new roman",12,"bold"),width=29)
        name_entry.grid(row=11,column=1,pady=5)
        
        #delievary
        delivery_lbl=Label(left_frame,text="Delivery:",font=("times new roman",12,"bold"))
        delivery_lbl.grid(row=12,column=0,sticky=W,pady=5)
        yes_btn1=Radiobutton(self.root,variable=self.delivery_var,value="Yes",text="Yes",font=("times new roman",13,"bold"))
        yes_btn1.place(x=130,y=460)
        self.delivery_var.set(value="No")
        
        no_btn1=Radiobutton(self.root,variable=self.delivery_var,value="No",text="No",font=("times new roman",13,"bold"))
        no_btn1.place(x=200,y=460)
#====================================left farme btn frame ========================================
        btn_frame=Frame(self.root,bd=3,relief=RIDGE)
        btn_frame.place(x=430,y=490,width=854,height=40)

        #add btn
        addbtn=Button(btn_frame,text="ADD",command=self.addcustomer,font=("times new roman",11,"bold"),width=22,bd=4,bg="green",fg="white",activebackground="green",activeforeground="white",cursor="hand2")
        addbtn.grid(row=0,column=0,)

        addbtn=Button(btn_frame,text="UPDATE",command=self.updatebtn,font=("times new roman",11,"bold"),width=22,bd=4,bg="orange",fg="white",activebackground="orange",activeforeground="white",cursor="hand2")
        addbtn.grid(row=0,column=1)

        addbtn=Button(btn_frame,text="DELETE",command=self.deletebtn,font=("times new roman",11,"bold"),width=22,bd=4,bg="red",fg="white",activebackground="red",activeforeground="white",cursor="hand2")
        addbtn.grid(row=0,column=2)

        addbtn=Button(btn_frame,text="RESET",command=self.reset_btn,font=("times new roman",11,"bold"),width=22,bd=4,bg="purple",fg="white",activebackground="purple",activeforeground="white",cursor="hand2")
        addbtn.grid(row=0,column=3)

        """img_pen_path=resource_path(r"E:\Tailor management system\photos\pen.png")
        img_pen=Image.open(img_pen_path)
        img_pen=img_pen.resize((45,330),Image.ANTIALIAS)
        self.pendemo=ImageTk.PhotoImage(img_pen)
        lbl_img_pen=Label(left_frame,image=self.pendemo,bd=0,relief=RIDGE)
        lbl_img_pen.place(x=363,y=15,height=330)"""

        img_pen_path=resource_path("photos\pen.png")
        img_pen=Image.open(img_pen_path)
        img_pen=img_pen.resize((45,330),Image.ANTIALIAS)
        self.pendemo=ImageTk.PhotoImage(img_pen)
        lbl_img_pen=Label(left_frame,image=self.pendemo,bd=0,relief=RIDGE)
        lbl_img_pen.place(x=363,y=15,height=330)


       
#==============================================right frame=========================================
        rightframe=Label(self.root,bd=10,relief=RIDGE)
        rightframe.place(x=430,y=55,width=867,height=435)
     

        """img_book_path=resource_path(r"E:\Tailor management system\photos\book.jpeg")
        img_book=Image.open(img_book_path)
        img_book=img_book.resize((848,120),Image.ANTIALIAS)
        self.bookdemo=ImageTk.PhotoImage(img_book)
        lbl_img_book=Label(rightframe,image=self.bookdemo,bd=0,relief=RIDGE,bg="black")
        lbl_img_book.place(x=0,y=0,height=120)"""

        img_book_path=resource_path(r"photos\book.jpeg")
        img_book=Image.open(img_book_path)
        img_book=img_book.resize((848,120),Image.ANTIALIAS)
        self.bookdemo=ImageTk.PhotoImage(img_book)
        lbl_img_book=Label(rightframe,image=self.bookdemo,bd=0,relief=RIDGE,bg="black")
        lbl_img_book.place(x=0,y=0,height=120)
#==============================search system=====================================
#==============================searching variable========================================
        self.searchcombobox=StringVar()
        self.searchentry=StringVar()
#================================================================================
        lbl_search=Label(rightframe,text="Search By",font=("times new roman",12,"bold"),bg="red",fg="white")
        lbl_search.place(x=0,y=123)

        search_combobox=ttk.Combobox(rightframe,textvariable=self.searchcombobox,font=("times new roman",12,"bold"),width=15,state="readonly")
        search_combobox["values"]=("Select Option","Reference_No","Name","Submitted_Date","Receiving_Date","Delivery")
        search_combobox.current(0)
        search_combobox.place(x=78,y=123)

        search_entry=Entry(rightframe,textvariable=self.searchentry,font=("times new roman",13,"bold"),width=15)
        search_entry.place(x=225,y=123)

        search_btn=Button(rightframe,command=self.search_btn,text="Search",font=("times new roman",10,"bold"),bg="black",fg="gold",width=15,activebackground="black",activeforeground="gold",cursor="hand2")
        search_btn.place(x=370,y=123)

        show_all_btn=Button(rightframe,command=self.fecth_data,text="Show All",font=("times new roman",10,"bold"),bg="black",fg="gold",width=15,activebackground="black",activeforeground="gold",cursor="hand2")
        show_all_btn.place(x=490,y=123)

        reset=Button(rightframe,command=self.reset_search_btn,text="Reset",font=("times new roman",10,"bold"),bg="black",fg="gold",width=15,activebackground="black",activeforeground="gold",cursor="hand2")
        reset.place(x=610,y=123)

        exportbtn=Button(rightframe,command=self.export_to_word,text="DOWNLOAD DATA",font=("times new roman",9,"bold"),bg="black",fg="gold",width=15,activebackground="black",activeforeground="gold",cursor="hand2")
        exportbtn.place(x=730,y=123)

#===========================================================scrolll bar==========================================
        customertable=Frame(rightframe,bd=4,relief=RIDGE)
        customertable.place(x=0,y=150,width=847,height=265)

        scroll_x2=Scrollbar(customertable,orient=HORIZONTAL)
        scroll_x2.pack(side=BOTTOM,fill=X)

        scroll_y2=Scrollbar(customertable,orient=VERTICAL)
        scroll_y2.pack(side=RIGHT,fill=Y)

        self.customertable=ttk.Treeview(customertable,
                                              columns=(
                                              "Reference_No",
                                              "Name",
                                              "Delivery",
                                              "Address",
                                              "Phone",
                                              "Shirt_Quantity",
                                              "Pant_Quantity",
                                              "Type_Of_Shirt",
                                              "Design_Of_Shirt",
                                              "Receiving_Date",
                                              "Submitted_Date",
                                              "Payment"),xscrollcommand=scroll_x2.set,yscrollcommand=scroll_y2.set)
        
        scroll_x2.pack(side=BOTTOM,fill=X)
        scroll_y2.pack(side=RIGHT,fill=Y)

        scroll_x2.config(command=self.customertable.xview)
        scroll_y2.config(command=self.customertable.yview)

        self.customertable["show"]="headings"

        self.customertable.heading("Reference_No",text="Reference_No")
        self.customertable.heading("Name",text="Name")
        self.customertable.heading("Delivery",text="Delivery")
        self.customertable.heading("Address",text="Address")
        self.customertable.heading("Phone",text="Phone")
        self.customertable.heading("Shirt_Quantity",text="Shirt_Quantity")
        self.customertable.heading("Pant_Quantity",text="Pant_Quantity")
        self.customertable.heading("Type_Of_Shirt",text="Type_Of_Shirt")
        self.customertable.heading("Design_Of_Shirt",text="Design_Of_Shirt")
        self.customertable.heading("Receiving_Date",text="Receiving_Date")
        self.customertable.heading("Submitted_Date",text="Submitted_Date")
        self.customertable.heading("Payment",text="Payment")
        self.customertable.pack(fill=BOTH,expand=1)

        self.customertable.column("Reference_No",width=100)
        self.customertable.column("Name",width=100)
        self.customertable.column("Delivery",width=100)
        self.customertable.column("Address",width=100)
        self.customertable.column("Phone",width=100)
        self.customertable.column("Shirt_Quantity",width=100)
        self.customertable.column("Pant_Quantity",width=100)
        self.customertable.column("Type_Of_Shirt",width=100)
        self.customertable.column("Design_Of_Shirt",width=100)
        self.customertable.column("Receiving_Date",width=100)
        self.customertable.column("Submitted_Date",width=100)
        self.customertable.column("Payment",width=100)
        self.fecth_data()
        self.customertable.bind("<ButtonRelease-1>",self.getcursor)
       
#============================================ database connectivity=============================================
        self.auto_export_to_csv()
    def addcustomer(self):
        if self.name.get()=="":
           messagebox.showerror("Error","All fields are required",parent=self.root)
        else:
          conn=get_connection()
          mycursor=conn.cursor()
          mycursor.execute("insert into addcustomer values(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)",(
                                                                                                self.reference_var.get(),
                                                                                                self.name.get(),
                                                                                                self.delivery_var.get(),
                                                                                                self.address.get(),
                                                                                                self.phone.get(),
                                                                                                self.Shirtquantity.get(),
                                                                                                self.Pantquantity.get(),
                                                                                                self.type_of_shirt.get(),
                                                                                                self.design_of_shirt.get(),
                                                                                                self.receiving_date.get(),
                                                                                                self.submitted_date.get(),
                                                                                                self.payment.get()
                                                                                                ))
        conn.commit()
        self.fecth_data()
        messagebox.showinfo("Success","Customer data has been successfully Added",parent=self.root)
        self.send_sms_add_button()
        self.getcursor()
        conn.close()
        

    def send_sms_add_button(self):
       
        # Ask confirmation before sending SMS
        ask = messagebox.askyesno("Send SMS", f"Do you want to send SMS to {self.name.get()}?", parent=self.root)
        if not ask:
            return

        try:
            # ---- Fetch URL from Database ----
            conn=get_connection()
            self.cursor=conn.cursor()
            self.cursor.execute("SELECT url FROM gatewayurl LIMIT 1")
            result = self.cursor.fetchone()

            if result is None:
                messagebox.showerror("Error", "❌ No Server URL found in database!", parent=self.root)
                return

            SIMPLE_GATEWAY_URL = result[0]     # URL loaded from DB

        except Exception as e:
            messagebox.showerror("Error", f"❌ Failed to fetch server URL:\n{e}", parent=self.root)
            return

        try:
            phone_number = self.phone.get()
            name = self.name.get()
            ref = self.reference_var.get()

            # SMS content
            message = (
                f"👋 Dear {name},\n"
                f"✅ Your order has been confirmed successfully.\n"
                f"🧾 Receipt No: {ref}\n"
                f"💫 Thank you for trusting Sarth Tailor Shop — "
                f"we value your style and look forward to your next visit! ✂️✨"
            )

            payload = {
                "phone": phone_number,
                "message": message
            }

            headers = {
                "Content-Type": "application/json"
            }

            # ---- Send SMS ----
            response = requests.post(SIMPLE_GATEWAY_URL, json=payload, headers=headers)
            response_text = response.text
            print("SMS Response:", response_text)

            # ---- Check success ----
            if response.status_code == 200 and ("success" in response_text or "SENT" in response_text):
                messagebox.showinfo("SMS Sent", "SMS has been sent successfully!", parent=self.root)
            else:
                messagebox.showwarning("SMS Failed", f"Failed to send SMS.\nServer Response: {response_text}", parent=self.root)

        except requests.exceptions.ConnectionError:
            messagebox.showwarning("Warning","Connection Error: Check Wi-Fi or ensure SMS Gateway is running.",parent=self.root)
        except Exception as e:
            print("SMS Error:", e)
            messagebox.showwarning("Warning",f"Customer added but SMS not sent.\nError: {e}",parent=self.root)




    def fecth_data(self):
        conn=get_connection()
        mycursor=conn.cursor()
        mycursor.execute("select * from addcustomer ORDER BY STR_TO_DATE(Receiving_Date,'%d-%m-%Y') DESC")
        rows=mycursor.fetchall()
        if len(rows)!=0:
                self.customertable.delete(*self.customertable.get_children())
                for i in rows:
                        self.customertable.insert("",END,values=i)
                conn.commit()
        conn.close()



    def getcursor(self,event=""):
        cursor=self.customertable.focus()
        content=self.customertable.item(cursor)
        rows=content["values"]

        self.reference_var.set(rows[0])
        self.name.set(rows[1]),
        self.delivery_var.set(rows[2]),
        self.address.set(rows[3]),
        self.phone.set(rows[4]),
        self.Shirtquantity.set(rows[5]),
        self.Pantquantity.set(rows[6]),
        self.type_of_shirt.set(rows[7]),
        self.design_of_shirt.set(rows[8]),
        self.receiving_date.set(rows[9]),
        self.submitted_date.set(rows[10]),
        self.payment.set(rows[11]),
        


    def updatebtn(self):
        conn=get_connection()
        mycursor=conn.cursor()
        mycursor.execute("update addcustomer set Name=%s,Delivery=%s,Address=%s,Phone=%s,Shirt_Quantity=%s,Pant_Quantity=%s,Type_Of_Shirt=%s,Design_Of_Shirt=%s,Receiving_Date=%s,Submitted_Date=%s,Payment=%s where Reference_No=%s",
                                                                                             (
                                                                                                self.name.get(),
                                                                                                self.delivery_var.get(),
                                                                                                self.address.get(),
                                                                                                self.phone.get(),
                                                                                                self.Shirtquantity.get(),
                                                                                                self.Pantquantity.get(),
                                                                                                self.type_of_shirt.get(),
                                                                                                self.design_of_shirt.get(),
                                                                                                self.receiving_date.get(),
                                                                                                self.submitted_date.get(),
                                                                                                self.payment.get(),
                                                                                                self.reference_var.get()
                                                                                                
                                                                                                
                                                                                                ))
        conn.commit()
        self.fecth_data()
        conn.close()
        messagebox.showinfo("Success","Data has been updated",parent=self.root)


    def deletebtn(self):
        mdelete=messagebox.askyesno("Sarth Tailor","Do you want delete this customer",parent=self.root)
        if mdelete>0:
            conn=get_connection()
            mycursor=conn.cursor()
            query="delete from addcustomer where Reference_No=%s"
            val=(self.reference_var.get(),)
            mycursor.execute(query,val)
        else:
            if not mdelete:
                return

        conn.commit()
        self.fecth_data()
        conn.close()
        messagebox.showinfo("Success","Data has been deleted",parent=self.root)

    def reset_btn(self):
        self.reference_var.set(random.randint(1000,9999))
        self.name.set("")
        self.delivery_var.set("No")
        self.address.set("")
        self.phone.set("")
        self.Shirtquantity.set("")
        self.Pantquantity.set("")
        self.type_of_shirt.set("Select Option")
        self.design_of_shirt.set("Select Option")
        self.receiving_date.set("")
        self.submitted_date.set("")
        self.payment.set("")
        

#================================================searching method===========================================
        
    def search_btn(self):
        if self.searchcombobox.get()=="Select Option" or self.searchentry.get()=="":
            messagebox.showerror("Error","All Fields Are Required",parent=self.root)
        else:
            conn=get_connection()
            mycursor=conn.cursor()
            mycursor.execute("SELECT * FROM addcustomer WHERE " + str(self.searchcombobox.get()) + " LIKE '%" + str(self.searchentry.get()) + "%'")
            rows=mycursor.fetchall()
            if len(rows)!=0:
                self.customertable.delete(*self.customertable.get_children())
                for i in rows:
                    self.customertable.insert("",END,values=i)
                conn.commit()
            conn.close()


    def reset_search_btn(self):
        self.searchcombobox.set("Select Option")
        self.searchentry.set("")


    def auto_export_to_csv(self):
        import csv
        import mysql.connector
        import hashlib

        try:
            conn = get_connection()
            cursor = conn.cursor()
            cursor.execute("SELECT * FROM addcustomer ORDER BY Reference_No ASC")
            all_data = cursor.fetchall()
            conn.close()

            # ---- Make a hash of current DB state ----
            db_hash = hashlib.md5(str(all_data).encode()).hexdigest()

            # ---- If previous hash exists, compare ----
            if hasattr(self, "last_db_hash") and self.last_db_hash == db_hash:
                print("Success","⏸ No change detected — CSV not updated.")
            else:
                self.last_db_hash = db_hash  # Save new hash

                headers = [
                    "Reference_No", "Name", "Delivery", "Address", "Phone",
                    "Shirt_Quantity", "Pant_Quantity", "Type_Of_Shirt",
                    "Design_Of_Shirt", "Receiving_Date", "Submitted_Date", "Payment"
                ]

                #file_path = r"E:\Tailor management system\addcustomer_data.csv"
                file_path =resource_path ("addcustomer_data.csv")

                with open(file_path, mode="w", newline="", encoding="utf-8") as file:
                    writer = csv.writer(file)
                    writer.writerow(headers)
                    writer.writerows(all_data)

                print("Success","CSV updated (change detected).")

        except Exception as e:
            print(f"Error updating CSV: {e}")

        # 🔁 Repeat every 5 seconds
        self.root.after(5000, self.auto_export_to_csv)


    def export_to_word(self):
        import pandas as pd
        from docx import Document
        from tkinter import filedialog, messagebox

        # Step 1: Confirmation popup
        confirm = messagebox.askyesno("Confirmation", "Do you really want to export customer data to a Word file?",parent=self.root)
        if not confirm:
            return  # Agar user "No" dabaye to function yahi ruk jaye

        try:
            #csv_path = "E:\\Tailor management system\\addcustomer_data.csv"
            csv_path = resource_path("addcustomer_data.csv")
            df = pd.read_csv(csv_path)

            #  Step 2: Save file dialog
            save_path = filedialog.asksaveasfilename(
                defaultextension=".docx",
                filetypes=[("Word Documents", "*.docx")],
                title="Save Word File As"
            )

            if not save_path:
                return  # User ne cancel kiya

            #  Step 3: Word document create karo
            doc = Document()
            doc.add_heading("Customer Details", level=1)

            #  Step 4: Table banao aur header add karo
            table = doc.add_table(rows=1, cols=len(df.columns))
            hdr_cells = table.rows[0].cells
            for i, col_name in enumerate(df.columns):
                hdr_cells[i].text = col_name

            #  Step 5: Data rows add karo
            for _, row in df.iterrows():
                row_cells = table.add_row().cells
                for i, value in enumerate(row):
                    row_cells[i].text = str(value)

            # Step 6: Save Word file
            doc.save(save_path)
            messagebox.showinfo("Success", f"Word file successfully saved at:\n{save_path}", parent=self.root)

        except Exception as e:
            messagebox.showerror("Error", f"Failed to export: {str(e)}", parent=self.root)



    

    


           



if __name__=="__main__":
    root=Tk()
    app=customer(root)
    root.mainloop()